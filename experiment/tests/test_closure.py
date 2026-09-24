"""闭环回归：不复制任务、候选筛选、源参数、压力基准和双精度Excel往返。"""
import unittest,tempfile,json,math,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from io_utils import read_csv
from closure_experiments import independent_charge
from q2_physics import charge_duration
from solve_q4 import FrozenQ3,strict_no_copy
from submission import save_workbook_lossless

class ClosureTests(unittest.TestCase):
    def test_excel_roundtrip_preserves_all_float_bits(self):
        from openpyxl import Workbook,load_workbook
        vals=[109.2088888888889,23.020833333333332,743.679230704014,math.nextafter(1.,2.),1e-20,0.,-98.,32.45]
        with tempfile.TemporaryDirectory() as d:
            wb=Workbook();ws=wb.active
            for x in vals:ws.append([x])
            p=Path(d)/'roundtrip.xlsx';save_workbook_lossless(wb,p);v=[r[0] for r in load_workbook(p,data_only=True).active.values]
            self.assertEqual([float(x).hex() for x in vals],[float(x).hex() for x in v])
    def test_nonfinite_excel_rejected(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            wb=Workbook();wb.active.append([float('nan')])
            with self.assertRaises(ValueError):save_workbook_lossless(wb,Path(d)/'bad.xlsx')
    def test_charge_independent_formula_dense(self):
        for i in range(1001):self.assertAlmostEqual(independent_charge(i/1000,1800),charge_duration(i/1000,1800),places=9)
    def test_closure_candidate_meets_three_components(self):
        self.assertGreaterEqual(strict_no_copy(FrozenQ3(ROOT))['component_count'],3)
    def test_no_relay_replica_in_q4(self):
        f=FrozenQ3(ROOT)
        for k in (2,3):
            r=read_csv(ROOT/f'results/q4/K{k}/中继任务_唯一执行.csv')
            self.assertEqual(len(r),len(f.relays))
    def test_q4_time_energy_same_as_frozen_q3(self):
        s=json.loads((ROOT/'results/q3/summary.json').read_text());q=json.loads((ROOT/'results/q4/summary.json').read_text())
        for k in ('2','3'):
            self.assertAlmostEqual(q['selected'][k]['total_energy_kwh'],s['total_energy_kwh'],places=9)
            self.assertAlmostEqual(q['selected'][k]['joint_makespan_s'],s['joint_makespan_s'],places=9)
    def test_energy_sweep_nominal_is_feasible(self):
        rr=read_csv(ROOT/'results/closure/fixed_energy_sweep.csv')
        for r in rr:
            if float(r['energy_multiplier'])==1:self.assertEqual(r['energy_and_battery_feasible'],'True')
    def test_charging_sweep_nominal_is_feasible(self):
        for r in read_csv(ROOT/'results/closure/fixed_charging_sweep.csv'):
            if float(r['full_charge_time_multiplier'])==1:self.assertEqual(r['feasible'],'True')
    def test_delay_zero_is_feasible(self):
        for r in read_csv(ROOT/'results/closure/common_delay_fine_sweep.csv'):
            if float(r['common_delay_s'])==0:self.assertEqual(r['hard_feasible'],'True')
    def test_delay_hard_violations_monotone(self):
        rr=read_csv(ROOT/'results/closure/common_delay_fine_sweep.csv')
        for q in ('Q2','Q3'):
            v=[int(r['hard_late_boxes']) for r in rr if r['question']==q];self.assertEqual(v,sorted(v))
    def test_selected_pool_rule(self):
        rr=[r for r in read_csv(ROOT/'results/q3/closure_existing_candidates.csv') if r['strict_K3_feasible']=='True']
        best=min(rr,key=lambda r:(float(r['weighted_tardiness_s']),float(r['joint_makespan_s']),float(r['total_energy_kwh']),int(r['transport_sorties'])+int(r['relay_sorties']),r['candidate']))
        c=json.loads((ROOT/'results/q3/closure_repair_certificate.json').read_text());self.assertEqual(c['selected_candidate']['candidate'],best['candidate'])
    def test_chinese_font_available(self):
        from paper_figures import cjk_font
        self.assertTrue(cjk_font())
if __name__=='__main__':unittest.main()
