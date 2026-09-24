"""结果模板同步与逐格验收。Python 3.13.5；工作簿制作使用openpyxl3.1.5。
标准库独立只读OOXML验证；每次从当次CSV完整重建主表及配套表，并从真实单元格反向验证。
"""
from pathlib import Path
from collections import Counter
import csv,json,math,shutil
from io_utils import xlsx_read,read_csv,write_csv,save_json,sha256
from q2_validation import independent_arcs,validate_rows
from config import CFG
INTEGER_HEADERS={'K（2或3）','A型运输无人机数','B型运输无人机数','C型运输无人机数','A型电池组数','B型电池组数','C型电池组数','中继无人机数','中继能源组件数'}
NUMERIC_HEADERS={'总质量（kg）','总体积（m³）','往返时间（s）','架次能耗（kWh）','返航SOC（%）','开始时刻（s）','返回O01时刻（s）','交付完成时刻（s）','结束时刻（s）','悬停经度（°）','悬停纬度（°）','悬停海拔（m）','建链完成时刻（s）','服务结束时刻（s）'} | INTEGER_HEADERS
SHEET_FILES={'Q1_单点组批':'results/q1/Q1_单点组批.csv','Q2_运输架次':'results/q2/Q2_运输架次.csv','Q2_逐箱交付':'results/q2/Q2_逐箱交付.csv','Q3_中继架次':'results/q3/Q3_中继架次.csv','Q3_通信保障':'results/q3/Q3_通信保障.csv','Q4_分区配置':'results/q4/Q4_分区配置.csv'}
CELL_REL_TOL=1e-13
CELL_ABS_TOL=1e-9
EXPORT_NAME='结果提交.xlsx'
COMPANION_NAME='Q3运输与逐箱补充.xlsx'
COMPANION_FILES={'Q3_运输架次':'results/q3/Q3_运输架次.csv','Q3_逐箱交付':'results/q3/Q3_逐箱交付.csv'}

def expected_matrices(root):
    root=Path(root);raw,_=xlsx_read(root/'data/raw/结果提交模板.xlsx');result={}
    for name,m in raw.items():
        headers=m[0][:]
        while headers and headers[-1] is None:headers.pop()
        if name in SHEET_FILES:
            with (root/SHEET_FILES[name]).open(encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f)
                if reader.fieldnames!=headers:raise AssertionError('结果CSV表头与源模板不一致：'+name)
                rows=list(reader)
            result[name]=[headers]+[[float(r[h]) if h in NUMERIC_HEADERS else r[h] for h in headers] for r in rows]
        else:result[name]=[headers]
    return result

def _compare(root,path,write_log=True,companion=False):
    wanted=companion_matrices(root) if companion else expected_matrices(root);actual,_=xlsx_read(path);checks=[];dicts={}
    def ck(n,ok):checks.append({'cell_or_check':n,'pass':bool(ok)})
    ck('sheet_names_and_order',list(actual)==list(wanted))
    for name,exp in wanted.items():
        mat=actual.get(name,[])
        # 允许原模板保留已格式化但完全空白的尾部行/列，不允许新增非空答案。
        while mat and not any(x is not None and x!='' for x in mat[-1]):mat.pop()
        ck(name+'/record_count',len(mat)==len(exp));cols=len(exp[0])
        for i,row in enumerate(exp):
            got=mat[i] if i<len(mat) else []
            ck(name+f'/row{i+1}_no_extra_fields',not any(x is not None and x!='' for x in got[cols:]))
            for j,val in enumerate(row):
                vv=got[j] if j<len(got) else None
                if isinstance(val,(int,float)):
                    ok=isinstance(vv,(int,float)) and math.isfinite(vv) and math.isclose(float(val),float(vv),rel_tol=CELL_REL_TOL,abs_tol=CELL_ABS_TOL)
                else:ok=val==vv or (val in (None,'') and vv in (None,''))
                ck(f'{name}!R{i+1}C{j+1}',ok)
        dicts[name]=[dict(zip(exp[0],row)) for row in mat[1:]]
    if write_log:write_csv(Path(root)/('results/logs/companion_cell_checks.csv' if companion else 'results/logs/submission_cell_checks.csv'),checks)
    return dicts,checks

def verify_submission(root,path=None,companion_path=None):
    root=Path(root);path=Path(path) if path else root/'submission'/EXPORT_NAME
    rows,cc=_compare(root,path)
    if not all(r['pass'] for r in cc):raise AssertionError('XLSX与当次CSV/原模板逐格核对失败')
    companion_path=Path(companion_path) if companion_path else root/'submission'/COMPANION_NAME
    q3rows,ccc=_compare(root,companion_path,companion=True)
    if not all(c['pass'] for c in ccc):raise AssertionError('Q3运输补充表与当次CSV不一致')
    data=json.loads((root/'data/cleaned/model_inputs.json').read_text('utf-8'));arc=independent_arcs(root,data)
    qs,ck,tr,bx,rs=validate_rows(data,arc,rows['Q2_运输架次'],rows['Q2_逐箱交付'])
    write_csv(root/'results/logs/workbook_physics_checks.csv',ck)
    # Q1从最终Excel重算，不以CSV相等替代物理验证。
    models={m['model_id']:m for m in data['transport_models']};boxes={b['box_id']:b for b in data['boxes']};seen=[];qc=[]
    for r in rows['Q1_单点组批']:
        sid=r['服务区编号'];m=models[r['机型编号']];ids=r['货箱编号列表'].split(';');seen+=ids
        bs=[boxes[b] for b in ids];w=sum(b['weight_kg'] for b in bs);v=sum(b['volume_m3'] for b in bs);E=0.;T=0.
        for (u,st,q) in [('O01',sid,w),(sid,'O01',0)]:
            a=arc[u,st];L=m['empty_range_m']-(m['empty_range_m']-m['full_range_m'])*(q/m['max_payload_kg'])**CFG.range_load_exponent
            E+=CFG.horizontal_energy_multiplier*m['usable_energy_kwh']*a['d']/L+CFG.climb_energy_multiplier*(m['empty_mass_kg']+q)*CFG.gravity_m_s2*a['up']/(m['climb_efficiency']*CFG.joules_per_kwh)
            T+=a['up']/m['climb_speed_mps']+a['d']/m['cruise_speed_mps']+a['down']/m['descent_speed_mps']
        T+=m['prepare_s']+m['load_per_box_s']*len(ids)+m['handoff_base_s']+m['handoff_per_box_s']*len(ids)
        vals={'mass':abs(w-float(r['总质量（kg）']))<1e-8,'volume':abs(v-float(r['总体积（m³）']))<1e-9,'energy':abs(E-float(r['架次能耗（kWh）']))<1e-8,
            'time_convention':abs(T-float(r['往返时间（s）']))<1e-6,'SOC':abs(100*(1-E/m['usable_energy_kwh'])-float(r['返航SOC（%）']))<1e-7,
            'capacity':w<=m['max_payload_kg']+1e-8 and v<=m['volume_m3']+1e-9,'reserve':1-E/m['usable_energy_kwh']>=m['reserve_fraction']-1e-9,
            'single_service':all(b['service_id']==sid for b in bs)}
        qc.extend({'check':str(r['架次编号'])+'/'+k,'pass':v} for k,v in vals.items())
    qc.append({'check':'Q1_exact_once_80_boxes','pass':Counter(seen)==Counter(boxes.keys())})
    write_csv(root/'results/logs/workbook_q1_physics_checks.csv',qc)
    from q3_validation import run_q3_validation
    v3=run_q3_validation(root,dense=False,trip_rows=q3rows['Q3_运输架次'],delivery_rows=q3rows['Q3_逐箱交付'],relay_rows=rows['Q3_中继架次'],comm_rows=rows['Q3_通信保障'],write=False)
    save_json(root/'results/logs/workbook_q3_independent_validation.json',v3)
    from q4_validation import validate_q4_template,IndependentQ4,verify_case
    iq4=IndependentQ4(root);q4ck=validate_q4_template(root,rows['Q4_分区配置'],iq4)
    q4cases={};q4lineage=[]
    for k in (2,3):
        lineage_checks,ps=verify_case(root,f'results/q4/K{k}',iq4,check_physics=True);q4lineage.extend(lineage_checks);q4cases[str(k)]=ps
    write_csv(root/'results/logs/workbook_q4_partition_checks.csv',q4ck)
    write_csv(root/'results/logs/workbook_q4_inheritance_resource_checks.csv',q4lineage)
    save_json(root/'results/logs/workbook_q4_physics_cases.json',q4cases)
    if any(not c['pass'] for c in q4ck+q4lineage) or any(c['checks_failed'] for c in q4cases.values()):
        raise AssertionError('最终Q4分区表、资源证书或固定物理任务复算未通过')
    summary={'q4_workbook_partition_checks':len(q4ck),'q4_workbook_partition_failed':sum(not c['pass'] for c in q4ck),
        'q4_workbook_inheritance_resource_checks':len(q4lineage),'q4_workbook_inheritance_failed':sum(not c['pass'] for c in q4lineage),
        'q4_workbook_physics_checks':sum(c['checks_total'] for c in q4cases.values()),'q4_workbook_physics_failed':sum(c['checks_failed'] for c in q4cases.values()),
        'q4_source_inventory_feasible':{k:c['source_inventory_sufficient'] for k,c in q4cases.items()},
        'q4_relay_replication_allowed':False,'q4_strict_unique_mission_enforced':True,
        'q3_workbook_physics_checks':v3['checks_total'],'q3_workbook_physics_failed':v3['checks_failed'],
        'companion_cell_checks':len(ccc),'companion_cell_checks_failed':sum(not c['pass'] for c in ccc),
        'companion_path':'submission/'+COMPANION_NAME,'companion_row_counts':{k:len(v) for k,v in q3rows.items()},
        'q3_recomputed_summary':v3,'xlsx_path':'submission/'+EXPORT_NAME,'xlsx_sha256':sha256(path),'original_sheet_names_and_headers_preserved':True,
        'sheet_row_counts':{k:len(v) for k,v in rows.items()},'cell_checks':len(cc),'cell_checks_failed':sum(not r['pass'] for r in cc),
        'q1_workbook_physics_checks':len(qc),'q1_workbook_physics_failed':sum(not r['pass'] for r in qc),
        'q2_workbook_physics_checks':qs['checks_total'],'q2_workbook_physics_failed':qs['checks_failed'],
        'q2_recomputed_summary':qs,'Q4_completed_5_resource_configuration_rows':True,'Q3_transport_in_separate_companion_preserves_Q2':True,'stored_numeric_values_not_display_rounding':True}
    save_json(root/'results/logs/submission_validation.json',summary)
    if qs['checks_failed'] or v3['checks_failed'] or any(not r['pass'] for r in qc):raise AssertionError('最终工作簿物理复核未通过')
    return summary

def companion_matrices(root):
    result={}
    for name,f in COMPANION_FILES.items():
        with (Path(root)/f).open(encoding='utf-8-sig',newline='') as h:
            reader=csv.DictReader(h);headers=reader.fieldnames;rows=list(reader)
        result[name]=[headers]+[[float(r[k]) if k in NUMERIC_HEADERS else r[k] for k in headers] for r in rows]
    return result


def sync_submission(root,force_rebuild=False):
    root=Path(root);out=root/'submission';out.mkdir(parents=True,exist_ok=True);dest=out/EXPORT_NAME;dest2=out/COMPANION_NAME
    dest.unlink(missing_ok=True);dest2.unlink(missing_ok=True)
    for name,f in (SHEET_FILES|COMPANION_FILES).items():shutil.copy2(root/f,out/(name+'.csv'))
    # 不再依赖可选artifact_tool或复制参考工作簿；始终由当次结果重建。
    build_openpyxl_workbook(root,dest,False)
    build_openpyxl_workbook(root,dest2,True)
    mode='both_rebuilt_openpyxl_from_current_results'
    summary=verify_submission(root,dest,dest2);save_json(out/'提交验收.json',summary)
    from q4_reports import write_submission_readme
    write_submission_readme(root)
    save_json(root/'results/logs/excel_sync_mode.json',{'mode':mode,'snapshot_is_not_solver_input':True,'guard_every_cell':True,'reject_stale_files':True,'Q2_not_overwritten_by_Q3':True})
    return summary


def build_openpyxl_workbook(root,destination,companion=False):
    """保留原模板名称/表头/顺序，完整浮点数存储；中文展示不改变计算值。"""
    from copy import copy
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment
    root=Path(root);mats=companion_matrices(root) if companion else expected_matrices(root)
    wb=Workbook() if companion else load_workbook(root/'data/raw/结果提交模板.xlsx')
    if companion:wb.remove(wb.active)
    for name,mat in mats.items():
        ws=wb.create_sheet(name) if companion else wb[name]
        if ws.max_row>1:ws.delete_rows(2,ws.max_row-1)
        for i,row in enumerate(mat,1):
            for j,v in enumerate(row,1):
                c=ws.cell(i,j,v);c.font=Font(name='微软雅黑',size=10,bold=(i==1))
                c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                if i==1:c.fill=PatternFill('solid',fgColor='D9E2F3')
                elif mat[0][j-1] in NUMERIC_HEADERS:c.number_format='0' if mat[0][j-1] in INTEGER_HEADERS else '0.000000'
            ws.row_dimensions[i].height=36 if i==1 else min(90,max(29,18*max((math.ceil(len(str(v or ''))/38) for v in row),default=1)))
        for j,h in enumerate(mat[0],1):
            ws.column_dimensions[get_column_letter(j)].width=48 if h in ('服务区列表','货箱编号列表') else (32 if h=='访问服务区顺序' else (24 if h=='货箱编号' else 19))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        if name=='Q1_单点组批':
            j=mat[0].index('往返时间（s）')+1;ws.cell(1,j).comment=Comment('本提交定义为准备、装载、往返飞行、交接的累计作业时间。纯飞行时间在结果台账另列。','计算口径')
        ws.sheet_view.showGridLines=False
    save_workbook_lossless(wb,destination)

def save_workbook_lossless(wb,destination):
    # openpyxl默认%.16g会截掉部分二进制浮点的往返精度；连续遮挡边界可能因此变侧。
    # 只为本次写入局部替换数值序列化，使用Python最短精确往返repr；不改任何计算阈值。
    import openpyxl.cell._writer as _writer
    old_safe=_writer.safe_string
    def exact_safe(value):
        if isinstance(value,float):
            if not math.isfinite(value):raise ValueError('不允许把非有限数值写入提交表')
            return repr(value)
        return old_safe(value)
    _writer.safe_string=exact_safe
    try:wb.save(destination)
    finally:_writer.safe_string=old_safe
    # 规范化仅用于容器元数据：实际执行时刻仍在results/logs记录。
    # 保证同一数据产生相同ZIP字节，不能由生成日期改变验收JSON中的SHA-256。
    import zipfile,xml.etree.ElementTree as ET
    destination=Path(destination)
    with zipfile.ZipFile(destination) as archive:payloads={n:archive.read(n) for n in archive.namelist()}
    if 'docProps/core.xml' in payloads:
        tree=ET.fromstring(payloads['docProps/core.xml'])
        for node in tree:
            if node.tag.split('}')[-1] in ('created','modified'):node.text='2000-01-01T00:00:00Z'
        payloads['docProps/core.xml']=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
        for name,body in sorted(payloads.items()):
            info=zipfile.ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o600<<16
            archive.writestr(info,body)

